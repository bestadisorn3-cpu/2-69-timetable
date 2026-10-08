import openpyxl,re,json,sys
"""เพิ่มรายชื่อครูของแผนก + คาบที่ไม่อยู่ในตารางเรียนของแผนก (เช่น สอนแผนกอื่น) จากไฟล์ตารางสอนรายบุคคล
ใช้: python build_teachers.py ตารางสอน.xlsx data.js"""
CODE=re.compile(r'^\d{5}-\d{4}$')
s=lambda v:'' if v is None else re.sub(r'\s+',' ',str(v)).strip()
wb=openpyxl.load_workbook(sys.argv[1]); out=sys.argv[2]
src=open(out).read(); T=json.loads(src[src.index('{'):src.rindex('}')+1])
SKIP={'ตรวจสอบยอดนักเรียน'}
teachers=[];log=[]
for ws in wb.worksheets:
  if ws.title in SKIP: continue
  key=ws.title.strip(); full=s(ws['P7'].value)
  if key=='ฝึกสอน': key='ฝึกสอน ('+re.sub(r'^(นาย|นางสาว|นาง)','',full).split(' ')[0]+')'
  subj={}
  for r in range(12,25):
    c=s(ws.cell(r,15).value)
    if CODE.match(c): subj[c]=s(ws.cell(r,16).value)
  mine=[(se['d'],se['c'],set(se['p'])) for c in T['classes'] for se in c['sessions'] if ws.title.strip() in se['t']]
  extra=[]
  for d in range(5):
    dr=6+4*d; cell=lambda rr,c: s(ws.cell(rr,c).value)
    segs=[];cur=None
    for c in range(4,15):
      cv=cell(dr,c)
      if CODE.match(cv):
        cur=dict(code=cv,cols=[c],rooms=[],name=cell(dr+1,c),label=cell(dr+3,c),open=True); segs.append(cur); continue
      if cur and cur['open']:
        cur['cols'].append(c)
        if cv: cur['rooms'].append(cv); cur['open']=False
        if not cur['label'] and cell(dr+3,c): cur['label']=cell(dr+3,c)
      elif cv or cell(dr+1,c):  # segment without code (name only)
        cur=dict(code='',cols=[c],rooms=[cv] if cv else [],name=cell(dr+1,c),label=cell(dr+3,c),open=not cv); segs.append(cur)
    for g in segs:
      lab=g['label']; lunch=8 if not lab.startswith('ส') else 7
      ps=[c-3 for c in g['cols'] if c!=lunch]
      code=g['code'] or next((k for k,v in subj.items() if g['name'].replace('ฯ','') in v),'')
      if any(d==md and code==mc and (set(ps)&mp) for md,mc,mp in mine): continue
      lab=re.sub(r'\s*\(\d+\)\s*$','',lab)
      extra.append(dict(d=d,c=code,p=ps,r=g['rooms'],n=subj.get(code,g['name']),l=lab))
      log.append((key,d,code,ps,lab,g['rooms']))
  teachers.append(dict(k=key,full=full,x=extra))
T['teachers']=teachers
open(out,'w').write('window.TT='+json.dumps(T,ensure_ascii=False,separators=(',',':'))+';\n')
for l in log: print('เพิ่ม',l)
print(len(teachers),'teachers')
