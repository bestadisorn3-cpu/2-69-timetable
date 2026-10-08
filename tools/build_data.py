import openpyxl,re,json,sys
wb=openpyxl.load_workbook(sys.argv[1])
CODE=re.compile(r'^\d{5}-\d{4}$')
def s(v): return '' if v is None else re.sub(r'\s+',' ',str(v)).strip()
TFIX={'อ.ธนณีย์':'อ.ธนนีย์'}
RFIX={'ทฐ2':'ทฐ 2','สปก':'สปก.'}
# ห้องที่ไม่มีใน Excel (ยืนยันโดยครู): (id ห้อง, วัน 0=จ.., รหัสวิชา) -> ห้อง
ROOM_OVERRIDE={('v1-5',1,'20000-1102'):'741',('v1-5',4,'20104-2013'):'825'}
ALIAS=[('วิทย์','วิทยาศาสตร์')]
def norm(x):
  x=x.replace('ฯ','').replace(' ','')
  for a,b in ALIAS: x=x.replace(a,b)
  return x
classes=[];issues=[]
for ws in wb.worksheets[1:]:
  heads=[c for row in ws.iter_rows() for c in row if isinstance(c.value,str) and 'ตารางชั้นเรียน' in c.value]
  for h in heads:
    H=h.row; title=s(h.value)
    m=re.search(r'(ปวช|ปวส)\.\s*(\d)/(\d+)',title)
    level,year,room=m.group(1),int(m.group(2)),int(m.group(3))
    g=re.search(r'กลุ่ม\s*(\d+)',title); group=int(g.group(1)) if g else None
    n=re.findall(r'\((\d+)\)',title); students=int(n[-1]) if n else None
    tags=[]
    if 'ม.6' in title: tags.append('ม.6')
    if 'ทวิ' in title or 'ทวิ' in ws.title: tags.append('ทวิภาคี')
    if 'กสศ' in title or 'กสศ' in ws.title: tags.append('กสศ.')
    lunch=next(c for c in range(4,15) if 'พักรับประทาน' in s(ws.cell(H+3,c).value))
    advisor=s(ws.cell(H+5,16).value)
    subj={};order=[]
    r=H+7
    while r<H+30:
      code=s(ws.cell(r,15).value)
      if code=='รวมทั้งสิ้น': break
      if CODE.match(code):
        subj[code]=dict(code=code,name=s(ws.cell(r,16).value),t=ws.cell(r,17).value or 0,p=ws.cell(r,18).value or 0,n=ws.cell(r,19).value or 0)
        order.append(code)
      r+=1
    segs=[]
    for d in range(5):
      dr=H+3+4*d
      cell=lambda rr,c: s(ws.cell(rr,c).value)
      cols=[c for c in range(4,15) if c!=lunch]
      cur=None
      def close():
        global cur
      day=[]
      for c in cols:
        cv,nv,tv=cell(dr,c),cell(dr+1,c),cell(dr+3,c)
        if CODE.match(cv):
          cur=dict(day=d,code=cv,names=[],cols=[],rooms=[],teachers=[],open=True); day.append(cur)
        elif not (cv or nv or tv):
          if cur and cur['open']: cur['cols'].append(c)
          continue
        elif cur is None or not cur['open']:
          cur=dict(day=d,code='',names=[],cols=[],rooms=[],teachers=[],open=True); day.append(cur)
        cur['cols'].append(c)
        if nv: cur['names'].append(nv)
        if tv: cur['teachers'].append(tv)
        if cv and not CODE.match(cv):
          cur['rooms'].append(cv); cur['open']=False
      # trim trailing empty cols of segments without room (no closing room)
      for sg in day:
        while sg['cols'] and not any(cell(dr+k,sg['cols'][-1]) for k in (0,1,3)): sg['cols'].pop()
      # merge codeless nameless segments
      out=[]
      for sg in day:
        if not sg['code'] and not sg['names'] and out:
          p=out[-1]
          if sg['rooms']: p['cols']+=sg['cols']; p['rooms']+=sg['rooms']
          p['teachers']+=sg['teachers']; continue
        out.append(sg)
      segs+=out
    # hours from coded
    cnt={}
    for sg in segs:
      if sg['code']: cnt[sg['code']]=cnt.get(sg['code'],0)+len(sg['cols'])
    for sg in segs:
      if sg['code']: continue
      nm=norm(sg['names'][0])
      cand=[k for k in order if norm(subj[k]['name']).startswith(nm) or nm in norm(subj[k]['name'])]
      if len(cand)>1:
        c2=[k for k in cand if subj[k]['t']+subj[k]['p']-cnt.get(k,0)>=len(sg['cols'])]
        if c2: cand=c2
      if len(cand)>=1:
        sg['code']=cand[0]; cnt[cand[0]]=cnt.get(cand[0],0)+len(sg['cols'])
        if len(cand)>1: issues.append(('AMBIG',title,sg['names'],cand))
      else: issues.append(('NOMATCH',title,sg['names']))
    sessions=[]
    for sg in segs:
      tch=[]
      for t in sg['teachers']:
        t=re.sub(r'^อ\.\s*','อ.',t); t=TFIX.get(t,t)
        if t not in tch: tch.append(t)
      sessions.append(dict(d=sg['day'],c=sg['code'],p=[c-3 for c in sg['cols']],r=list(dict.fromkeys(RFIX.get(x,x) for x in sg['rooms'])),t=tch,s=sg['names'][0] if sg['names'] else ''))
    for k in order:
      if cnt.get(k,0)!=subj[k]['t']+subj[k]['p']: issues.append(('HRS',title,k,subj[k]['name'],subj[k]['t']+subj[k]['p'],cnt.get(k,0)))
    for se in sessions:
      if se['c'] not in subj: issues.append(('NOSUBJ',title,se))
      if not se['t']: issues.append(('NOTEACH',title,se['c'],se['d']))
      if not se['r']: issues.append(('NOROOM',title,se['c'],se['d']))
    cid=f"{'v' if level=='ปวช' else 's'}{year}-{room}"+(f"-g{group}" if group else '')
    for se in sessions:
      k=(cid,se['d'],se['c'])
      if k in ROOM_OVERRIDE and not se['r']: se['r']=[ROOM_OVERRIDE[k]]
    classes.append(dict(id=cid,level=level+'.',year=year,room=room,group=group,students=students,tags=tags,advisor=advisor,lunch=lunch-3,subjects=[subj[k] for k in order],sessions=sessions))
for i in issues: print(i)
mx=max(max(p for se in c['sessions'] for p in se['p']) for c in classes); print('max period',mx)
data=dict(term='2/2569',college='วิทยาลัยเทคนิคปราจีนบุรี',dept='แผนกวิชาช่างไฟฟ้ากำลัง',classes=classes)
open(sys.argv[2],'w').write('window.TT='+json.dumps(data,ensure_ascii=False,separators=(',',':'))+';\n')
print(len(classes),'classes')
teachers=sorted({t for c in classes for se in c['sessions'] for t in se['t']}); print(len(teachers),teachers)
rooms=sorted({r for c in classes for se in c['sessions'] for r in se['r']}); print(rooms)
